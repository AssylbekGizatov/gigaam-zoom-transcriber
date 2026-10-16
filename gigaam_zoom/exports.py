"""Explicit local transcript exports, without network calls or overwriting files.

TXT, Markdown and JSON iterate the SQLite store. PDF paginates each line as it
arrives (ReportLab still retains its PDF page objects). python-docx retains the
document tree, so large meetings should use TXT or JSON on small computers.
"""
from __future__ import annotations

from dataclasses import asdict, is_dataclass
from datetime import datetime
import json
import math
from pathlib import Path
import re
import sys
from typing import Any, Iterable

EXPORT_FORMATS = ("TXT", "MD", "DOCX", "PDF", "JSON")
TITLE = "Транскрипция встречи Zoom"
TIMING_NOTE = (
    "Метки реплик приблизительные. "
    "Участник определяется по идентификатору источника Zoom; общий микрофон не разделяется на людей. "
    "Предварительный текст может измениться. Пропуски аудио перечислены отдельно."
)
_SENSITIVE = re.compile(r"secret|token|password|authorization|credential|signature|passcode|join_url|meeting_url", re.I)
_SEGMENT_FIELDS = ("id", "speaker_id", "display_name", "start", "end", "text", "status", "source_epoch")


def format_time(seconds: float) -> str:
    if not math.isfinite(float(seconds)):
        raise ValueError("Некорректная метка времени")
    value = max(0, round(float(seconds) * 1000))
    hours, value = divmod(value, 3_600_000)
    minutes, value = divmod(value, 60_000)
    whole, millis = divmod(value, 1000)
    return f"{hours:02d}:{minutes:02d}:{whole:02d}.{millis:03d}"


def _mapping(value: Any) -> dict:
    if isinstance(value, dict):
        return value
    if is_dataclass(value):
        return asdict(value)
    try:
        return dict(value)
    except (TypeError, ValueError):
        return {key: getattr(value, key) for key in _SEGMENT_FIELDS if hasattr(value, key)}


def _safe_metadata(value: Any) -> Any:
    """Allow JSON values and remove authentication fields, including nested ones.

    This is defense in depth: the store must never contain credentials in the
    first place. Recognized speech itself is user data and is not redacted.
    """
    if isinstance(value, dict):
        return {str(k): _safe_metadata(v) for k, v in value.items() if not _SENSITIVE.search(str(k))}
    if isinstance(value, (list, tuple)):
        return [_safe_metadata(item) for item in value]
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    return str(value)


def _record(segment: Any) -> dict:
    values = _mapping(segment)
    result = {key: values.get(key) for key in _SEGMENT_FIELDS}
    result["display_name"] = str(result["display_name"] or "Неизвестный участник")
    result["speaker_id"] = str(result["speaker_id"] or "unknown")
    result["status"] = str(result["status"] or "final")
    result["text"] = str(result["text"] or "")
    for key in ("start", "end"):
        result[key] = float(result[key] or 0)
        if not math.isfinite(result[key]) or result[key] < 0:
            raise ValueError("Некорректные метки времени в хранилище")
    if result["end"] < result["start"]:
        raise ValueError("Конец реплики предшествует началу")
    return result


def _label(record: dict) -> str:
    return (f"[{format_time(record['start'])} - {format_time(record['end'])}] "
            f"{record['display_name']} [ID {record['speaker_id']}; "
            f"epoch {record['source_epoch']}; {record['status']}; segment {record['id']}]")


def _events(store: Any) -> Iterable[dict]:
    if hasattr(store, "iter_events"):
        for event in store.iter_events():
            yield _safe_metadata(_mapping(event))
    elif hasattr(store, "iter_gaps"):
        for event in store.iter_gaps():
            yield {"kind": "audio_gap", **_safe_metadata(_mapping(event))}


def _timing_note(metadata: dict) -> str:
    if metadata.get("time_basis") == "meeting":
        origin = "Время отсчитывается от подтверждённого начала встречи Zoom. "
    else:
        origin = "Время отсчитывается от начала локально получаемой сессии RTMS, а не от начала всей встречи. "
    return origin + TIMING_NOTE


def _md(value: str) -> str:
    return re.sub(r"([\\`*_{}\[\]<>#!|])", r"\\\1", value)


def _text(store: Any, stream: Any, markdown: bool, metadata: dict) -> None:
    escape = _md if markdown else str
    stream.write(("# " if markdown else "") + TITLE + "\n\n" + _timing_note(metadata) + "\n\n")
    if metadata:
        stream.write(escape("Метаданные: " + json.dumps(metadata, ensure_ascii=False)) + "\n\n")
    for value in store.iter_segments():
        record = _record(value)
        label = escape(_label(record))
        stream.write((f"**{label}**" if markdown else label) + "\n\n")
        stream.write(escape(record["text"]) + "\n\n")
    stream.write(("## " if markdown else "") + "События и пропуски\n\n")
    found = False
    for event in _events(store):
        found = True
        stream.write(escape(json.dumps(event, ensure_ascii=False)) + "\n\n")
    if not found:
        stream.write("Нет зарегистрированных событий.\n")


def _json(store: Any, stream: Any, metadata: dict) -> None:
    stream.write('{"schema_version":1,"timing":')
    json.dump("meeting_seconds" if metadata.get("time_basis") == "meeting" else "rtms_session_seconds", stream)
    stream.write(',"note":')
    json.dump(_timing_note(metadata), stream, ensure_ascii=False)
    stream.write(',"metadata":')
    json.dump(metadata, stream, ensure_ascii=False, allow_nan=False)
    stream.write(',"segments":[')
    first = True
    for value in store.iter_segments():
        if not first:
            stream.write(",")
        first = False
        json.dump(_record(value), stream, ensure_ascii=False, allow_nan=False)
    stream.write('],"events":[')
    first = True
    for event in _events(store):
        if not first:
            stream.write(",")
        first = False
        json.dump(event, stream, ensure_ascii=False, allow_nan=False)
    stream.write("]}\n")


def _font_path() -> Path:
    roots = [Path(__file__).resolve().parents[1]]
    if getattr(sys, "frozen", False):
        roots.insert(0, Path(sys.executable).resolve().parent)
    if getattr(sys, "_MEIPASS", None):
        roots.insert(0, Path(sys._MEIPASS))
    for root in roots:
        for path in (root / "fonts" / "DejaVuSans.ttf", root / "gigaam_zoom" / "fonts" / "DejaVuSans.ttf"):
            if path.is_file():
                return path
    raise FileNotFoundError("Не найден локальный шрифт fonts/DejaVuSans.ttf для PDF")


def _lines(store: Any, metadata: dict) -> Iterable[tuple[str, int]]:
    yield TITLE, 17
    yield _timing_note(metadata), 9
    if metadata:
        yield "Метаданные: " + json.dumps(metadata, ensure_ascii=False), 9
    for value in store.iter_segments():
        record = _record(value)
        yield _label(record), 9
        yield record["text"], 11
    yield "События и пропуски", 14
    found = False
    for event in _events(store):
        found = True
        yield json.dumps(event, ensure_ascii=False), 9
    if not found:
        yield "Нет зарегистрированных событий.", 10


def _wrapped(text: str, width: float, font: str, size: int) -> Iterable[str]:
    """Wrap incrementally, including unbroken IDs and Unicode strings."""
    from reportlab.pdfbase import pdfmetrics
    for paragraph in text.replace("\r", "").split("\n"):
        line = ""
        for match in re.finditer(r"\S+|\s+", paragraph):
            token = match.group()
            if line and pdfmetrics.stringWidth(line + token, font, size) > width:
                yield line.rstrip()
                line = ""
            if not line:
                token = token.lstrip()
            # Very long source IDs must still fit the physical page width.
            for char in token:
                if line and pdfmetrics.stringWidth(line + char, font, size) > width:
                    yield line
                    line = ""
                line += char
        yield line.rstrip()


def _pdf(store: Any, stream: Any, metadata: dict) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen.canvas import Canvas
    font = "GigaAMZoomDejaVu"
    if font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font, str(_font_path())))
    canvas = Canvas(stream, pagesize=A4, pageCompression=1)
    canvas.setTitle(TITLE)
    canvas.setAuthor("")
    canvas.setCreator("GigaAM Zoom Transcriber")
    width, height = A4
    margin, bottom, y, page = 48, 48, height - 48, 1

    def footer():
        canvas.setFont(font, 8)
        canvas.drawRightString(width - margin, 25, str(page))

    for paragraph, size in _lines(store, metadata):
        if paragraph.startswith("[") and y - 60 < bottom:
            footer()
            canvas.showPage()
            page += 1
            y = height - margin
        for line in _wrapped(paragraph, width - 2 * margin, font, size):
            if y - size * 1.4 < bottom:
                footer()
                canvas.showPage()
                page += 1
                y = height - margin
            canvas.setFont(font, size)
            canvas.drawString(margin, y, line)
            y -= size * 1.4
        y -= 8
    footer()
    canvas.save()


def _docx(store: Any, stream: Any, metadata: dict) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    title = document.styles["Title"]
    title.font.name = "DejaVu Sans"
    title.font.size = Pt(20)
    title.font.color.rgb = RGBColor(0, 0, 0)
    document.core_properties.title = TITLE
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    first = True
    for text, size in _lines(store, metadata):
        if first:
            document.add_paragraph(text, style="Title")
            first = False
        else:
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.keep_with_next = size in (14,) or text.startswith("[")
            for run in paragraph.runs:
                run.font.size = Pt(size)
                run.bold = size == 14
    document.save(stream)


def _destination(path: Path | str) -> Path:
    text = str(path)
    if text.startswith(("\\\\", "//")) or "://" in text:
        raise ValueError("Выберите локальную папку компьютера")
    target = Path(path).expanduser().absolute()
    if not target.parent.is_dir():
        raise FileNotFoundError("Папка экспорта не существует")
    from .config import local_directory
    local_directory(str(target.parent))
    # Reject a symlinked network destination too, when detectable by the OS.
    resolved = str(target.parent.resolve())
    if resolved.startswith(("\\\\", "//")):
        raise ValueError("Сетевая папка экспорта не поддерживается")
    return target


def export_transcript(store: Any, path: Path | str, format: str) -> Path:
    """Create a selected local file; failures remove only this new partial file."""
    kind = format.upper().lstrip(".")
    if kind not in EXPORT_FORMATS:
        raise ValueError("Поддерживаются TXT, MD, JSON, DOCX и PDF")
    target = _destination(path)
    metadata = _safe_metadata(store.get_metadata())
    created = False
    try:
        if kind in ("TXT", "MD", "JSON"):
            with target.open("x", encoding="utf-8", newline="\n") as stream:
                created = True
                if kind == "JSON":
                    _json(store, stream, metadata)
                else:
                    _text(store, stream, kind == "MD", metadata)
        else:
            with target.open("xb") as stream:
                created = True
                (_pdf if kind == "PDF" else _docx)(store, stream, metadata)
    except BaseException:
        if created:
            target.unlink(missing_ok=True)
        raise
    return target


def default_export_name(format: str) -> str:
    return "Транскрипция_" + datetime.now().strftime("%Y-%m-%d_%H-%M-%S_%f") + "." + format.lower().lstrip(".")
